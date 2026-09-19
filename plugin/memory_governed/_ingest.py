# -*- coding: utf-8 -*-
"""Ingestion helpers for the governed knowledge base.

把「外部资料 → 结构化文本」的机械活收在这里，供 agent 的摄入工具调用。

**职责边界**（重要）：本模块只负责**抓取/读取/转录**，不做「归纳总结」。
归纳由 agent 的 LLM 完成（决策：归纳用 agent 同款 LLM），落库用
``governed_kb_add``。这样任何 agent 都能走同一条链路：

    ingest 工具（拿内容） → agent 归纳 → governed_kb_add（落库）

三类摄入，全部**可选依赖降级**，绝不抛异常（返回 ``{ok, error}``）：

- ``fetch_url``: URL → 正文（HTML 提取 / 纯文本直读）
- ``read_file``: 本地文件 → 文本（txt/md/pdf/docx；pdf/docx 依赖可选）
- ``transcribe_audio``: 音频 → 转写文本（ASR API，OpenAI 兼容
  ``/audio/transcriptions``，默认硅基流动 XingChenASR）
- ``_transcribe_chat_audio``: 另一种风格的兄弟原语 —— chat-completions 端点，音频
  以 base64 ``input_audio`` 传入（如小米 MiMo）；由 ``config.asr.api_style`` 选择
- ``transcribe_audio_auto``: 上面的封装，**长录音自动拆段**再逐段转写并拼接，
  结果按段合并（部分失败时保留已拿到的文本）—— 所有工具面都调它。它通过
  ``_transcribe_one`` 按 ``asr.api_style`` 分派，故两种风格共享拆段/归一化机制。

统一返回结构 ``{ok: bool, ..., error?: str}``，方便工具层直接透传。
"""

from __future__ import annotations

import base64
import html as _html
import logging
import math
import os
import re as _re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

#: 摄入内容的最大字符数（超长截断，避免爆 agent context；归纳用原文更优，
#: 但不可控的网页/文件需要上限）。
MAX_CONTENT_CHARS: int = 50000

_TAG_RE = _re.compile(r"<[^>]+>")
#: 正文无关块（脚本/样式/导航/页眉页脚），连同内容整体删除
_NOISE_BLOCK_RE = _re.compile(
    r"<(script|style|nav|header|footer|aside|noscript|form)[^>]*>.*?</\1>",
    _re.S | _re.I,
)
#: 块级标签 → 换行（保留段落结构）
_BLOCK_RE = _re.compile(
    r"</?(p|div|li|h[1-6]|tr|br|section|article|blockquote|pre)[^>]*>",
    _re.I,
)
_TITLE_RE = _re.compile(r"<title[^>]*>(.*?)</title>", _re.S | _re.I)

_AUDIO_MIME = {
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".m4a": "audio/mp4",
    ".ogg": "audio/ogg",
    ".flac": "audio/flac",
    ".webm": "audio/webm",
    ".mp4": "audio/mp4",
}


# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------

def _http_get(url: str, timeout: float) -> tuple[int, str, Dict[str, str]]:
    """GET 抓取。优先 httpx，降级 urllib（标准库）。返回 (status, text, headers)。"""
    headers = {"User-Agent": "hermes-memory-governed/0.1 (+knowledge-base)"}
    try:
        import httpx  # type: ignore

        r = httpx.get(url, timeout=timeout, follow_redirects=True, headers=headers)
        return r.status_code, r.text, {k.lower(): v for k, v in r.headers.items()}
    except ImportError:
        pass
    import urllib.error
    import urllib.request

    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:  # noqa: S310
            body = resp.read().decode("utf-8", errors="replace")
            return resp.status, body, {k.lower(): v for k, v in resp.headers.items()}
    except urllib.error.HTTPError as e:
        return e.code, "", {}
    except Exception as e:  # noqa: BLE001
        logger.debug("http get failed for %s: %s", url, e)
        return 0, "", {}


def _html_to_text(html_str: str) -> str:
    """HTML → 纯文本（内置轻量提取器，无第三方依赖）。"""
    text = _NOISE_BLOCK_RE.sub(" ", html_str)
    text = _BLOCK_RE.sub("\n", text)
    text = _TAG_RE.sub(" ", text)
    text = _html.unescape(text)
    text = _re.sub(r"[ \t\r]+", " ", text)
    text = _re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _truncate(content: str, limit: int = MAX_CONTENT_CHARS) -> tuple[str, bool]:
    if len(content) <= limit:
        return content, False
    return content[:limit], True


def _is_html(content_type: str) -> bool:
    return "html" in (content_type or "").lower()


def _read_text(path: Path) -> str:
    """读文本文件，utf-8 → gbk → 兜底 errors=replace。"""
    for enc in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


# ---------------------------------------------------------------------------
# 摄入：URL
# ---------------------------------------------------------------------------

def fetch_url(url: str, timeout: float = 15.0) -> Dict[str, Any]:
    """抓取 URL 并提取正文。返回 ``{ok, url, title, content, content_type, truncated}``。"""
    url = (url or "").strip()
    if not url:
        return {"ok": False, "error": "Missing url"}
    if not url.lower().startswith(("http://", "https://")):
        return {"ok": False, "error": f"Unsupported scheme (http/https only): {url}"}

    try:
        status, body, headers = _http_get(url, timeout)
    except Exception as e:  # noqa: BLE001 — 摄入工具绝不抛给 agent
        return {"ok": False, "error": f"fetch failed: {e}"}
    if status == 0:
        return {"ok": False, "error": f"fetch failed (network error): {url}"}
    if status >= 400:
        return {"ok": False, "error": f"HTTP {status}: {url}"}

    content_type = headers.get("content-type", "")
    if _is_html(content_type):
        title_m = _TITLE_RE.search(body)
        title = _html.unescape(title_m.group(1)).strip() if title_m else url.rstrip("/").split("/")[-1] or url
        content = _html_to_text(body)
    else:
        # 纯文本 / JSON / CSV 等直接返回原文（JSON 保持可读）
        title = url.rstrip("/").split("/")[-1] or url
        content = body

    if not content.strip():
        return {"ok": False, "error": f"no extractable content from {url}"}
    content, truncated = _truncate(content)
    return {
        "ok": True,
        "url": url,
        "title": title,
        "content": content,
        "content_type": content_type.split(";")[0].strip() or "text",
        "truncated": truncated,
    }


# ---------------------------------------------------------------------------
# 摄入：本地文件
# ---------------------------------------------------------------------------

_TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".json", ".csv", ".tsv", ".log",
    ".py", ".js", ".ts", ".html", ".htm", ".xml", ".yaml", ".yml",
    ".rst", ".tex", ".org",
}


def _read_pdf(path: Path) -> Dict[str, Any]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return {"ok": False, "error": "reading PDF requires pypdf: pip install pypdf"}
    try:
        reader = PdfReader(str(path))
        parts = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(parts)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": f"pdf extract failed: {e}"}
    if not text.strip():
        return {"ok": False, "error": f"no extractable text (scanned PDF?): {path.name}"}
    return {"ok": True, "content": text}


def _read_docx(path: Path) -> Dict[str, Any]:
    try:
        import docx  # type: ignore
    except ImportError:
        return {"ok": False, "error": "reading docx requires python-docx: pip install python-docx"}
    try:
        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        # 表格也纳入
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(" | ".join(cells))
        text = "\n".join(paras)
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": f"docx extract failed: {e}"}
    if not text.strip():
        return {"ok": False, "error": f"no extractable text: {path.name}"}
    return {"ok": True, "content": text}


def read_file(path: str) -> Dict[str, Any]:
    """读本地文件提取文本。按扩展名分派（txt/md 直读、pdf/docx 可选依赖）。

    返回 ``{ok, path, title, content, content_type, truncated}``。
    """
    p = Path((path or "").strip())
    if not p.exists():
        return {"ok": False, "error": f"file not found: {path}"}
    if not p.is_file():
        return {"ok": False, "error": f"not a file: {path}"}

    suffix = p.suffix.lower()
    title = p.stem

    if suffix in _TEXT_SUFFIXES:
        try:
            content = _read_text(p)
        except OSError as e:
            return {"ok": False, "error": f"read failed: {e}"}
    elif suffix == ".pdf":
        r = _read_pdf(p)
        if not r.get("ok"):
            return r
        content = r["content"]
    elif suffix == ".docx":
        r = _read_docx(p)
        if not r.get("ok"):
            return r
        content = r["content"]
    else:
        return {"ok": False, "error": f"unsupported file type: {suffix or '(none)'}"}

    if not content.strip():
        return {"ok": False, "error": f"empty content: {path}"}
    content, truncated = _truncate(content)
    return {
        "ok": True,
        "path": str(p),
        "title": title,
        "content": content,
        "content_type": suffix.lstrip(".") or "text",
        "truncated": truncated,
    }


# ---------------------------------------------------------------------------
# 摄入：音频转写（ASR）
# ---------------------------------------------------------------------------

def transcribe_audio(path: str, config: Any, timeout: float = 180.0) -> Dict[str, Any]:
    """音频 → 转写文本（OpenAI 兼容 ``/audio/transcriptions``）。

    ASR 配置在 ``config.asr``（provider/api_key_env/base_url/model/language）。
    依赖 httpx（multipart 上传）；未装或未配置时返回清晰的 ``error``。
    """
    p = Path((path or "").strip())
    if not p.exists():
        return {"ok": False, "error": f"file not found: {path}"}

    asr = getattr(config, "asr", None)
    base_url = str(getattr(asr, "base_url", "") or "").strip()
    model = str(getattr(asr, "model", "") or "").strip()
    if not base_url or not model:
        return {"ok": False, "error": "ASR not configured (asr.base_url / asr.model)"}

    api_key_env = str(getattr(asr, "api_key_env", "") or "")
    api_key = os.environ.get(api_key_env, "") if api_key_env else ""
    if not api_key:
        return {"ok": False, "error": f"ASR api key env not set: {api_key_env}"}

    try:
        import httpx  # type: ignore
    except ImportError:
        return {"ok": False, "error": "transcription requires httpx: pip install httpx"}

    mime = _AUDIO_MIME.get(p.suffix.lower(), "application/octet-stream")
    endpoint = base_url.rstrip("/") + "/audio/transcriptions"
    data: Dict[str, Any] = {"model": model}
    language = str(getattr(asr, "language", "") or "").strip()
    if language:
        data["language"] = language

    try:
        with p.open("rb") as f:
            audio_bytes = f.read()
        r = httpx.post(
            endpoint,
            headers={"Authorization": f"Bearer {api_key}"},
            files={"file": (p.name, audio_bytes, mime)},
            data=data,
            timeout=timeout,
        )
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": f"transcription request failed: {e}"}

    if r.status_code >= 400:
        return {"ok": False, "error": f"ASR HTTP {r.status_code}: {r.text[:300]}"}
    try:
        payload = r.json()
        text = payload.get("text", "") or ""
    except Exception:  # noqa: BLE001
        return {"ok": False, "error": "ASR returned non-JSON response"}

    if not text.strip():
        return {"ok": False, "error": "ASR returned empty transcript"}
    text, truncated = _truncate(text)
    return {
        "ok": True,
        "path": str(p),
        "text": text,
        "model": model,
        "truncated": truncated,
    }


def _transcribe_chat_audio(path: str, config: Any, timeout: float) -> Dict[str, Any]:
    """音频 → 转写文本（chat-completions + ``input_audio``，如小米 MiMo）。

    与 :func:`transcribe_audio` 的关键差别：这类端点**不是** OpenAI audio 兼容的，
    ``/audio/transcriptions`` 直接 404。转写走 ``/chat/completions``，音频以 base64
    **data URL** 放进 user 消息的 ``input_audio`` 内容块，文本在
    ``choices[0].message.content``（没有顶层 ``text``）。language 放在**顶层**
    ``asr_options``，留空表示自动识别。

    任何输入都**先整体转成 mp3**（无条件，复用 :func:`_transcode_to_mp3`）：mp3 是
    端点接受的两种格式之一，且 64 kbps 约 8 KB/s，base64 后远低于
    ``asr.max_encoded_bytes`` —— 3 分钟 wav 原始 5.6MB → base64 7.5MB，会顶到 10MB
    上限。这也让任何 ffmpeg 能解码的输入都能用这个风格。
    """
    p = Path((path or "").strip())
    if not p.exists():
        return {"ok": False, "error": f"file not found: {path}"}

    asr = getattr(config, "asr", None)
    base_url = str(getattr(asr, "base_url", "") or "").strip()
    model = str(getattr(asr, "model", "") or "").strip()
    if not base_url or not model:
        return {"ok": False, "error": "ASR not configured (asr.base_url / asr.model)"}

    api_key_env = str(getattr(asr, "api_key_env", "") or "")
    api_key = os.environ.get(api_key_env, "") if api_key_env else ""
    if not api_key:
        return {"ok": False, "error": f"ASR api key env not set: {api_key_env}"}

    try:
        import httpx  # type: ignore
    except ImportError:
        return {"ok": False, "error": "transcription requires httpx: pip install httpx"}

    tmp_dir = tempfile.mkdtemp(prefix="hgm_chat_")
    try:
        # 无条件转 mp3（端点的两种可接受格式之一，且体积可控）
        try:
            mp3_path = _transcode_to_mp3(p, tmp_dir)
        except Exception as e:  # noqa: BLE001
            return {
                "ok": False,
                "error": (
                    f"cannot decode '{p.suffix or '(none)'}' audio "
                    f"(unsupported by this ffmpeg build?): {e}"),
            }

        try:
            encoded = base64.b64encode(mp3_path.read_bytes())
        except OSError as e:
            return {"ok": False, "error": f"read failed: {e}"}

        max_bytes = _positive_int(getattr(asr, "max_encoded_bytes", None), 10_000_000)
        if len(encoded) > max_bytes:
            # 显式报错，不静默截断、不自动改 chunk —— 这个项目要的是明确拒绝
            return {
                "ok": False,
                "error": (
                    f"encoded audio is {len(encoded)} bytes, over the "
                    f"asr.max_encoded_bytes limit of {max_bytes} bytes; "
                    f"lower asr.chunk_minutes to split it into smaller pieces"),
            }

        body: Dict[str, Any] = {
            "model": model,
            "messages": [{
                "role": "user",
                "content": [{
                    "type": "input_audio",
                    "input_audio": {"data": "data:audio/mpeg;base64," + encoded.decode("ascii")},
                }],
            }],
        }
        language = str(getattr(asr, "language", "") or "").strip()
        if language:
            # 只在校验非空时给 asr_options —— 空值即自动识别，且不是每个
            # chat-completions 后端都能容忍未知的顶层键
            body["asr_options"] = {"language": language}

        endpoint = base_url.rstrip("/") + "/chat/completions"
        try:
            r = httpx.post(
                endpoint,
                headers={"Authorization": f"Bearer {api_key}",
                         "Content-Type": "application/json"},
                json=body,
                timeout=timeout,
            )
        except Exception as e:  # noqa: BLE001
            return {"ok": False, "error": f"transcription request failed: {e}"}

        if r.status_code >= 400:
            return {"ok": False, "error": f"ASR HTTP {r.status_code}: {r.text[:300]}"}
        try:
            payload = r.json()
        except Exception:  # noqa: BLE001
            return {"ok": False, "error": "ASR returned non-JSON response"}

        choices = payload.get("choices") or []
        if not choices:
            return {
                "ok": False,
                "error": ("ASR returned no choices (top-level keys: "
                          f"{sorted(payload.keys())})"),
            }
        message = choices[0].get("message") or {}
        text = str(message.get("content") or "")
        if not text.strip():
            # 兜底：某些实现把文本放在 choices[0].text
            text = str(choices[0].get("text") or "")
        if not text.strip():
            # 报出**实际看到的形状**，而不是笼统的「no text」
            shape = (sorted(choices[0].keys()) if isinstance(choices[0], dict)
                     else repr(choices[0])[:120])
            return {
                "ok": False,
                "error": f"ASR returned empty transcript (choices[0] keys: {shape})",
            }

        text, truncated = _truncate(text)
        return {
            "ok": True,
            "path": str(p),
            "text": text,
            "model": model,
            "truncated": truncated,
        }
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _transcribe_one(path: str, config: Any, timeout: float) -> Dict[str, Any]:
    """按 ``config.asr.api_style`` 把一次转写分派到对应的原语。

    唯一的风格分派点：``transcribe_audio_auto`` 的单次分支和逐段循环都走这里，
    于是拆段 / 归一化 / 部分失败这些机制对两种风格是共享的，而不是各写一份。
    """
    asr = getattr(config, "asr", None)
    style = str(getattr(asr, "api_style", "") or "").strip()
    if style == "chat_audio":
        return _transcribe_chat_audio(path, config, timeout)
    return transcribe_audio(path, config, timeout=timeout)


# ---------------------------------------------------------------------------
# 摄入：音频转写（长录音自动拆段）
# ---------------------------------------------------------------------------

#: 外部工具名 → 环境变量覆盖名。两者都在 PATH 缺失时可用。
_FFMPEG_ENV = {"ffmpeg": "HGM_FFMPEG", "ffprobe": "HGM_FFPROBE"}


def _ffmpeg_tool(name: str) -> Optional[str]:
    """解析外部媒体工具路径：环境变量覆盖优先，否则查 PATH。

    ``name`` 为 ``"ffmpeg"`` 或 ``"ffprobe"``。找不到时返回 ``None`` —— 调用方
    据此给出「缺 ffmpeg」而非「超时」的清晰错误。
    """
    env_key = _FFMPEG_ENV.get(name, "")
    if env_key:
        override = (os.environ.get(env_key) or "").strip()
        if override:
            return override
    return shutil.which(name)


def probe_audio_duration(path: str) -> Optional[float]:
    """尽力探测音频时长（秒）。真正未知时返回 ``None``，绝不抛异常。

    优先 ``ffprobe``；ffprobe 缺失或失败时，对 ``.wav`` 退回标准库 ``wave``。
    """
    p = Path((path or "").strip())
    if not p.exists():
        return None

    ffprobe = _ffmpeg_tool("ffprobe")
    if ffprobe:
        try:
            proc = subprocess.run(
                [ffprobe, "-v", "error", "-show_entries", "format=duration",
                 "-of", "default=nw=1:nk=1", str(p)],
                capture_output=True, text=True, timeout=30,
            )
            if proc.returncode == 0:
                for line in reversed((proc.stdout or "").splitlines()):
                    line = line.strip()
                    if not line or line.lower() == "n/a":
                        continue
                    try:
                        dur = float(line)
                    except ValueError:
                        break
                    if math.isfinite(dur) and dur >= 0:
                        return dur
                    break
        except Exception as e:  # noqa: BLE001 — 探测失败不是错误，只是未知
            logger.debug("ffprobe failed for %s: %s", p, e)

    # 标准库兜底：仅 .wav 可读
    if p.suffix.lower() == ".wav":
        try:
            import wave

            with wave.open(str(p), "rb") as w:
                frames = w.getnframes()
                rate = w.getframerate()
            if rate:
                return frames / float(rate)
        except Exception as e:  # noqa: BLE001
            logger.debug("wave fallback failed for %s: %s", p, e)

    return None


def split_audio(path: str, chunk_seconds: float, out_dir: Any) -> List[Path]:
    """把音频切成 ``chunk_seconds`` 秒的 mp3 片段，返回播放顺序的路径列表。

    每段都用 ffmpeg **重新编码为 mp3**（``-c:a libmp3lame -b:a 64k``）。这一步
    是刻意的、承重的：它把任何 ffmpeg 能解码的输入归一化成 ASR 端点接受的格式。
    （``.silk`` 不在其列 —— 本机 ffmpeg 既无 silk demuxer 也无 silk 解码器，无法
    支持；不受支持的容器在 :func:`transcribe_audio_auto` 里显式报错，绝不上传原文。）

    仅当 ffmpeg 本身失败时抛 ``RuntimeError``（调用方会转成 error 信封）。
    """
    ffmpeg = _ffmpeg_tool("ffmpeg")
    if not ffmpeg:
        raise RuntimeError(
            "ffmpeg not found: required to split long audio "
            "(install ffmpeg or set HGM_FFMPEG)")

    src = Path((path or "").strip())
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    if chunk_seconds <= 0:
        raise RuntimeError("chunk_seconds must be positive")
    duration = probe_audio_duration(str(src))
    if duration is None:
        raise RuntimeError(f"cannot determine audio duration for {src.name}")

    count = max(1, int(math.ceil(duration / chunk_seconds)))
    chunks: List[Path] = []
    for i in range(count):
        start = i * chunk_seconds
        dst = out / f"chunk_{i:04d}.mp3"
        try:
            proc = subprocess.run(
                [ffmpeg, "-hide_banner", "-loglevel", "error", "-nostdin",
                 "-ss", f"{start:.3f}", "-i", str(src),
                 "-t", f"{chunk_seconds:.3f}", "-vn",
                 "-c:a", "libmp3lame", "-b:a", "64k", "-y", str(dst)],
                capture_output=True, text=True, timeout=600,
            )
        except Exception as e:  # noqa: BLE001
            raise RuntimeError(f"ffmpeg could not run: {e}") from e
        if proc.returncode != 0 or not dst.exists():
            detail = (proc.stderr or "").strip().splitlines()
            tail = detail[-1] if detail else "unknown error"
            raise RuntimeError(
                f"ffmpeg failed splitting {src.name} chunk {i}: {tail[:200]}")
        chunks.append(dst)
    return chunks


def _transcode_to_mp3(src: Any, out_dir: Any) -> Path:
    """把任意 ffmpeg 可解码的输入**整体**转成一个 mp3。

    用于把 :data:`_AUDIO_MIME` 之外的容器（如 ``.amr``）归一化成 ASR 端点接受的
    格式。之所以必须做：端点**根本没有 amr 解码器** —— 原文上传（不论
    ``audio/amr`` / ``audio/amr-nb`` / ``application/octet-stream``）一律 HTTP 400。
    归一化与长度无关，所以短录音同样要走这里（旧的实现只在拆段时才归一化，于是
    几秒钟的微信/QQ 语音反而必失败）。

    仅当 ffmpeg 缺失或失败时抛 ``RuntimeError``。
    """
    ffmpeg = _ffmpeg_tool("ffmpeg")
    if not ffmpeg:
        raise RuntimeError(
            "ffmpeg not found: required to normalise this container "
            "(install ffmpeg or set HGM_FFMPEG)")

    source = Path(str(src))
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    dst = out / (source.stem + ".mp3")
    try:
        proc = subprocess.run(
            [ffmpeg, "-hide_banner", "-loglevel", "error", "-nostdin",
             "-i", str(source), "-vn",
             "-c:a", "libmp3lame", "-b:a", "64k", "-y", str(dst)],
            capture_output=True, text=True, timeout=600,
        )
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"ffmpeg could not run: {e}") from e
    if proc.returncode != 0 or not dst.exists():
        detail = (proc.stderr or "").strip().splitlines()
        tail = detail[-1] if detail else "unknown error"
        raise RuntimeError(
            f"ffmpeg failed transcoding {source.name}: {tail[:200]}")
    return dst


def _positive_float(value: Any, default: float) -> float:
    """转成正浮点，否则退回 ``default``；绝不抛异常。"""
    try:
        num = float(value)
    except (TypeError, ValueError):
        return default
    if not math.isfinite(num) or num <= 0:
        return default
    return num


def _positive_int(value: Any, default: int) -> int:
    """转成正整数，否则退回 ``default``；绝不抛异常。"""
    try:
        num = int(value)
    except (TypeError, ValueError):
        return default
    if num <= 0:
        return default
    return num


def transcribe_audio_auto(path: str, config: Any) -> Dict[str, Any]:
    """音频 → 转写文本，**容器归一化 + 长录音自动拆段**。所有工具面的统一入口。

    两步，顺序固定：

    1. **容器归一化**：``_AUDIO_MIME`` 之外的容器（如 ``.amr``）无论长短都先整体
       转成 mp3 —— ASR 端点根本无法解码它们，原文上传只会 HTTP 400。短录音同样
       要归一化（微信/QQ 语音就是几秒的 amr，那才是最常见、也是过去唯一必失败
       的情形）。ffmpeg 缺失或无法解码该容器时返回结构化错误，绝不上传原文。
    2. **单次 or 拆段**：时长在 ``config.asr.chunk_minutes`` 分钟以内（或探测不到）
       时单次请求；更长时切成 ``ceil(duration / (chunk_minutes*60))`` 段 mp3，
       逐段转写后按顺序用 ``"\\n\\n"`` 拼接，最后过一遍 :func:`_truncate`。

    全部成功**且确实产出了文本**才返回 ``ok: True``。任一段失败则绝不报成功：返回
    ``ok: False`` + ``partial: True`` + ``chunks_failed``（**1-based**），并保留已
    拿到的文本 —— 因为第 4 段碰到 502 就丢掉整场 50 分钟的会议是不可接受的，而假装
    成功更糟。
    """
    asr = getattr(config, "asr", None)
    timeout = _positive_float(getattr(asr, "timeout_seconds", None), 600.0)
    chunk_minutes = _positive_float(getattr(asr, "chunk_minutes", None), 10.0)
    threshold = chunk_minutes * 60.0

    # 空路径：明确报错，而不是退化成对空串的存在性/权限错误
    raw_path = (path or "").strip()
    if not raw_path:
        return {"ok": False, "error": "no path given"}

    src = Path(raw_path)
    suffix = src.suffix.lower()

    # 临时目录从一开始就建，finally 里每次都清 —— 归一化和拆段共用它，任何提前
    # 返回（包括错误分支）都不会漏掉清理。
    tmp_dir = tempfile.mkdtemp(prefix="hgm_asr_")
    try:
        work_path = raw_path
        if suffix not in _AUDIO_MIME:
            # 端点解不了的容器：先整体归一化成 mp3（与长度无关）
            if _ffmpeg_tool("ffmpeg") is None:
                return {
                    "ok": False,
                    "path": str(src),
                    "error": (
                        f"unsupported audio container '{suffix or '(none)'}' and "
                        f"ffmpeg is unavailable to normalise it; install ffmpeg or "
                        f"set HGM_FFMPEG. Accepted containers: "
                        f"{', '.join(sorted(_AUDIO_MIME))}"),
                }
            try:
                work_path = str(_transcode_to_mp3(src, tmp_dir))
            except Exception as e:  # noqa: BLE001
                return {
                    "ok": False,
                    "path": str(src),
                    "error": (
                        f"cannot decode '{suffix or '(none)'}' audio "
                        f"(unsupported by this ffmpeg build?): {e}"),
                }

        try:
            duration = probe_audio_duration(work_path)
        except Exception as e:  # noqa: BLE001 — 探测绝不该变成一次失败
            logger.debug("duration probe raised for %s: %s", work_path, e)
            duration = None

        # 时长未知或足够短：单次请求，原样返回（保持历史行为）
        if duration is None or duration <= threshold:
            result = _transcribe_one(work_path, config, timeout)
            # 归一化时转写的是临时 mp3，而那个临时目录马上就会在 finally 里被删。
            # 把它的路径当作 path 返回，对一个要落库/回读的 agent 毫无意义 —— 信封里
            # 的 path 永远指调用方传入的那个文件（长录音分支也是这么做的）。
            if work_path != raw_path and isinstance(result, dict):
                result["path"] = str(src)
            return result

        # 长录音：必须有 ffmpeg 才能拆 —— 缺了就说清楚缺什么，而不是等到超时
        if _ffmpeg_tool("ffmpeg") is None:
            return {
                "ok": False,
                "path": str(src),
                "error": (
                    f"audio is {duration:.0f}s but the split threshold is "
                    f"{threshold:.0f}s and ffmpeg is unavailable to split it; "
                    f"install ffmpeg or set HGM_FFMPEG"),
            }

        try:
            chunks = split_audio(work_path, threshold, tmp_dir)
        except Exception as e:  # noqa: BLE001
            return {"ok": False, "path": str(src),
                    "error": f"audio split failed: {e}"}

        # 拆段却一个片段都没有：绝不报成功（一个没有内容的成功报告正是本项目
        # 反复被坑的那类静默失败）
        if not chunks:
            return {
                "ok": False,
                "path": str(src),
                "text": "",
                "chunks": 0,
                "duration_seconds": round(duration, 1),
                "error": "audio split produced no chunks",
            }

        texts: List[str] = []
        failed: List[tuple] = []
        for idx, chunk in enumerate(chunks):
            r = _transcribe_one(str(chunk), config, timeout)
            if r.get("ok"):
                piece = str(r.get("text") or "").strip()
                if piece:
                    texts.append(piece)
                    continue
            failed.append((idx, str(r.get("error") or "unknown error")))

        text, truncated = _truncate("\n\n".join(texts))
        result: Dict[str, Any] = {
            # 成功必须同时满足：没有失败段，且确实拿到了文本
            "ok": (not failed) and bool(texts),
            "path": str(src),
            "text": text,
            "model": str(getattr(asr, "model", "") or ""),
            "truncated": truncated,
            "chunks": len(chunks),
            "duration_seconds": round(duration, 1),
        }
        if failed:
            total = len(chunks)
            result["partial"] = True
            # 1-based：chunks_failed:[3] 挨着 chunks:7 会被读成「七个里的第三个」，
            # 而它其实是第 4 个；字段只给人和读错误的 agent 看，必须人类可读。
            result["chunks_failed"] = [i + 1 for i, _ in failed]
            result["error"] = "; ".join(
                f"chunk {i + 1}/{total} failed: {err}" for i, err in failed)
        return result
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
